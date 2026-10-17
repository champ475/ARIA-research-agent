"""OPS-05: the capacity model's arithmetic, its workbook, and the limits CAPACITY.md names."""

from __future__ import annotations

import copy
from pathlib import Path

import pytest

from scripts import capacity_model as cm

ROOT = Path(__file__).resolve().parents[1]


def test_cost_per_turn_follows_tokens_and_price():
    i = cm.DEFAULTS
    # equity: 3309 in, 700 out at $0.15 / $0.60 per million, 88 INR per dollar
    assert cm.cost_per_turn_inr(i, "equity") == pytest.approx((3309 * 0.15 + 700 * 0.60) / 1e6 * 88)
    assert (
        cm.cost_per_turn_inr(i, "report")
        > cm.cost_per_turn_inr(i, "equity")
        > cm.cost_per_turn_inr(i, "alert")
    )
    assert sum(cm.MIX.values()) == pytest.approx(1.0)


def test_the_first_three_limits_and_their_thresholds():
    rows = cm.limits(cm.DEFAULTS)
    assert [r["limit"].split(" ")[0] for r in rows] == ["Database", "LLM", "App"]  # soonest first
    assert [round(r["dau"]) for r in rows] == [85, 639, 6912]
    assert rows[0]["dau"] < rows[1]["dau"] < rows[2]["dau"]


def test_a_second_replica_and_more_keys_move_their_own_limits_only():
    base = {r["limit"].split(" ")[0]: r["dau"] for r in cm.limits(cm.DEFAULTS)}
    i = copy.deepcopy(cm.DEFAULTS)
    i.replicas, i.llm_keys = 2, 5
    now = {r["limit"].split(" ")[0]: r["dau"] for r in cm.limits(i)}
    assert now["App"] == pytest.approx(base["App"] * 2)
    assert now["LLM"] == pytest.approx(base["LLM"] * 5)
    assert now["Database"] == pytest.approx(base["Database"])


def test_more_turns_per_user_brings_every_limit_closer():
    heavy = copy.deepcopy(cm.DEFAULTS)
    heavy.turns_per_dau_per_day = 10
    for a, b in zip(cm.limits(cm.DEFAULTS), cm.limits(heavy), strict=False):
        assert b["dau"] == pytest.approx(a["dau"] / 2) or b["limit"] != a["limit"]
    assert all(r["dau"] > 0 for r in cm.limits(heavy))


def test_trial_token_cap_note():
    assert "60,000-token daily cap" in cm.trial_token_cap_note(cm.DEFAULTS)


def test_every_input_has_a_status_and_a_source():
    for name, (status, source) in cm.NOTES.items():
        assert status in {"measured", "assumed", "verify"} and len(source) > 10, name
    scalars = {f for f in cm.Inputs.__dataclass_fields__} - {"calls_per_turn", "tokens_in", "tokens_out"}
    assert scalars <= set(cm.NOTES), scalars - set(cm.NOTES)


def test_workbook_has_the_inputs_and_live_formulas(tmp_path):
    from openpyxl import load_workbook

    path = tmp_path / "m.xlsx"
    cm.build_workbook(cm.DEFAULTS, path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Inputs", "Model"]
    inputs = {row[0].value: row[1].value for row in wb["Inputs"].iter_rows(min_row=2) if row[0].value}
    assert inputs["db_limit_mb"] == 512 and inputs["process_turns_per_second"] == 2.4
    statuses = {row[0].value: row[2].value for row in wb["Inputs"].iter_rows(min_row=2) if row[0].value}
    assert statuses["llm_requests_per_key_per_day"] == "verify"
    assert (
        statuses["turns_per_dau_per_day"] == "assumed" and statuses["process_turns_per_second"] == "measured"
    )
    formulas = [
        c.value
        for row in wb["Model"].iter_rows()
        for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    ]
    assert (
        len(formulas) >= 20
        and any("SUMPRODUCT" in f for f in formulas)
        and any("SMALL" in f for f in formulas)
    )
    # nothing in the Model sheet is a hard-coded number: it all moves with the Inputs
    assert not [
        c.value for row in wb["Model"].iter_rows(min_row=2) for c in row if isinstance(c.value, int | float)
    ]


def test_the_committed_workbook_is_the_current_model(tmp_path):
    from openpyxl import load_workbook

    fresh = tmp_path / "fresh.xlsx"
    cm.build_workbook(cm.DEFAULTS, fresh)

    def dump(path):
        wb = load_workbook(path)
        return {ws.title: [[c.value for c in row] for row in ws.iter_rows()] for ws in wb}

    assert dump(ROOT / "docs" / "capacity_model.xlsx") == dump(fresh), (
        "re-run: python scripts/capacity_model.py"
    )


def test_capacity_doc_quotes_the_model():
    doc = (ROOT / "docs" / "CAPACITY.md").read_text(encoding="utf-8")
    for r in cm.limits(cm.DEFAULTS):
        assert f"{r['dau']:,.0f}" in doc, r["limit"]
    assert f"₹{sum(cm.MIX[k] * cm.cost_per_turn_inr(cm.DEFAULTS, k) for k in cm.KINDS):.2f}" in doc


def test_summary_table_lists_every_kind_and_limit():
    text = cm.summary(cm.DEFAULTS)
    for kind in cm.KINDS:
        assert f"| {kind} |" in text
    assert text.count("| 1 |") == 1 and "weighted by the mix" in text
